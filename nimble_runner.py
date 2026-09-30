#!/usr/bin/env python3
"""
Nimble plant batch classifier (v1.0-dev).

One control workbook; choose engine nimble (local Ollama) or jev (TypeSafe).
Writes engine-prefixed sheets (output_nimble_* / output_jev_*) and never deletes
the other engine's results. When both Results sheets exist, refreshes output_compare.

Usage
    python nimble_runner.py workbook.xlsx --engine nimble
    python nimble_runner.py workbook.xlsx --engine jev --test
    python nimble_runner.py workbook.xlsx --validate-only

Requires: pip install requests openpyxl
  - nimble: Ollama >= 0.35 with `ollama pull nimble`
  - jev:    export TYPESAFE_API_KEY=...
"""
import argparse
import hashlib
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ID_RE = re.compile(r"^[A-Za-z0-9_\-]+$")
REQUIRED_SHEETS = ["Config", "Questions", "Labels", "Records"]
ENGINES = {
    "nimble": {
        "model": "nimble",
        "base_url": "http://localhost:11434",
        "endpoint": "/v1/systemone",
    },
    "jev": {
        "model": "jev-1.13.0",
        "base_url": "https://api.typesafe.ai",
        "endpoint": "/v1/systemone",
    },
}
DEFAULTS = {
    "engine": "nimble",
    "model": "nimble",
    "base_url": "http://localhost:11434",
    "endpoint": "/v1/systemone",
    "timeout_seconds": 120,
    "retries": 2,
    "default_review_below": 0.6,
    "test_first_n": 10,
    "record_text_column": "text",
    "output_prefix": "nimble_results",
    "api_key": "",
}
LEGACY_OUTPUT_SHEETS = ("output_Results", "output_Raw", "output_Summary", "output_Run_info")
OUTPUT_SUFFIXES = ("Results", "Raw", "Summary", "Run_info")


# ----------------------------------------------------------------- helpers
def s(v):
    """Cell value -> stripped string ('' for empty)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_table(ws):
    """Rows of a sheet as dicts keyed by header. Skips fully blank rows."""
    rows = ws.iter_rows(values_only=True)
    headers = [s(h) for h in next(rows, [])]
    out = []
    for n, row in enumerate(rows, start=2):
        rec = {h: v for h, v in zip(headers, row) if h}
        if all(s(v) == "" for v in rec.values()):
            continue
        rec["_row"] = n
        out.append(rec)
    return out


def output_sheet_names(engine):
    return tuple(f"output_{engine}_{suf}" for suf in OUTPUT_SUFFIXES)


def resolve_engine(cfg, cli_engine=None):
    eng = s(cli_engine if cli_engine else cfg.get("engine")).lower() or "nimble"
    if eng not in ENGINES:
        sys.exit(f"Unknown engine '{eng}'. Use one of: {', '.join(sorted(ENGINES))}")
    return eng


def apply_engine_preset(cfg, engine):
    """Apply built-in model/base_url/endpoint for the selected engine."""
    preset = ENGINES[engine]
    cfg["engine"] = engine
    cfg["model"] = preset["model"]
    cfg["base_url"] = preset["base_url"]
    cfg["endpoint"] = preset["endpoint"]
    return cfg


# ----------------------------------------------------------------- workbook
def load_spec(path):
    wb = load_workbook(path, data_only=True)
    missing = [n for n in REQUIRED_SHEETS if n not in wb.sheetnames]
    if missing:
        sys.exit(f"Workbook is missing sheet(s): {', '.join(missing)}")

    cfg = dict(DEFAULTS)
    for r in read_table(wb["Config"]):
        if s(r.get("key")) and r.get("value") is not None and s(r.get("value")) != "":
            cfg[s(r["key"])] = r["value"]
    cfg["timeout_seconds"] = to_float(cfg["timeout_seconds"]) or 120
    cfg["retries"] = int(to_float(cfg["retries"]) or 0)
    cfg["default_review_below"] = to_float(cfg["default_review_below"]) or 0.6
    cfg["test_first_n"] = int(to_float(cfg["test_first_n"]) or 0)

    questions = []
    for r in read_table(wb["Questions"]):
        if not s(r.get("question_id")):
            continue
        questions.append({
            "id": s(r["question_id"]),
            "equipment": s(r.get("equipment_type")),
            "type": s(r.get("type")).lower(),
            "instructions": s(r.get("instructions")),
            "active": s(r.get("active")).upper() != "N",
            "review_below": to_float(r.get("review_below")),
            "_row": r["_row"],
        })

    labels = {}
    for r in read_table(wb["Labels"]):
        qid, lab = s(r.get("question_id")), s(r.get("label"))
        if not qid and not lab:
            continue
        order = to_float(r.get("order"))
        labels.setdefault(qid, []).append(
            {"order": order if order is not None else 1e9 + r["_row"],
             "label": lab, "desc": s(r.get("description")), "_row": r["_row"]})
    for qid in labels:
        labels[qid].sort(key=lambda x: x["order"])

    text_col = s(cfg["record_text_column"])
    records = []
    for r in read_table(wb["Records"]):
        expected = {k[len("expected_"):]: s(v).lower() for k, v in r.items()
                    if k.startswith("expected_") and s(v) != ""}
        meta = {k: s(v) for k, v in r.items()
                if k not in ("_row", text_col, "check") and not k.startswith("expected_")}
        records.append({
            "id": s(r.get("record_id")), "equipment": s(r.get("equipment_type")),
            "text": s(r.get(text_col)), "expected": expected, "meta": meta, "_row": r["_row"],
        })
    return {"cfg": cfg, "questions": questions, "labels": labels, "records": records}


def validate(spec):
    errs = []
    seen = set()
    qids = {q["id"] for q in spec["questions"]}
    for q in spec["questions"]:
        where = f"Questions row {q['_row']} ({q['id']})"
        if not ID_RE.match(q["id"]):
            errs.append(f"{where}: question_id may only use letters, digits, _ and -")
        if q["id"] in seen:
            errs.append(f"{where}: duplicate question_id")
        seen.add(q["id"])
        if not q["instructions"]:
            errs.append(f"{where}: missing instructions")
        n = len(spec["labels"].get(q["id"], []))
        if q["type"] not in ("choice", "noul", "score"):
            errs.append(f"{where}: type must be choice, noul or score")
        elif q["type"] == "noul" and n:
            errs.append(f"{where}: noul questions take no labels (found {n})")
        elif q["type"] in ("choice", "score") and n < 2:
            errs.append(f"{where}: needs at least 2 labels (found {n})")
        if q["review_below"] is not None and not 0 <= q["review_below"] <= 1:
            errs.append(f"{where}: review_below must be between 0 and 1")
    for qid, items in spec["labels"].items():
        if qid not in qids:
            errs.append(f"Labels: unknown question_id '{qid}'")
        names = set()
        for it in items:
            if not it["label"]:
                errs.append(f"Labels row {it['_row']}: missing label")
            elif not ID_RE.match(it["label"]):
                errs.append(f"Labels row {it['_row']}: label '{it['label']}' may only use letters, digits, _ and -")
            elif it["label"] in names:
                errs.append(f"Labels row {it['_row']}: duplicate label '{it['label']}' for {qid}")
            names.add(it["label"])
    ids = set()
    for r in spec["records"]:
        if not r["id"]:
            errs.append(f"Records row {r['_row']}: missing record_id")
        elif r["id"] in ids:
            errs.append(f"Records row {r['_row']}: duplicate record_id '{r['id']}'")
        ids.add(r["id"])
        if not r["text"]:
            errs.append(f"Records row {r['_row']} ({r['id']}): empty text")
    if not spec["records"]:
        errs.append("Records sheet has no rows")
    if not any(q["active"] for q in spec["questions"]):
        errs.append("No active questions")
    return errs


def questions_hash(spec):
    payload = json.dumps({"q": [{k: v for k, v in q.items() if k != "_row"} for q in spec["questions"]],
                          "l": {k: [{a: b for a, b in i.items() if a != "_row"} for i in v]
                                for k, v in sorted(spec["labels"].items())}}, sort_keys=True)
    return hashlib.sha256(payload.encode()).hexdigest()[:12]


def questions_for(record, spec):
    rec_eq = record["equipment"].lower()
    return [q for q in spec["questions"]
            if q["active"] and (not q["equipment"] or q["equipment"].lower() == rec_eq)]


def build_request(record, qs, spec):
    qdict = {}
    for q in qs:
        d = {"type": q["type"], "instructions": q["instructions"]}
        labs = spec["labels"].get(q["id"], [])
        if q["type"] == "choice":
            d["criteria"] = {i["label"]: (i["desc"] or None) for i in labs}
        elif q["type"] == "score":
            d["criteria"] = [i["desc"] or i["label"] for i in labs]
        qdict[q["id"]] = d
    return {"model": spec["cfg"]["model"], "state": record["text"], "questions": qdict}


# ----------------------------------------------------------------- model calls
def resolve_api_key(cfg):
    return s(os.environ.get("TYPESAFE_API_KEY")) or s(cfg.get("api_key"))


def is_local_ollama(base_url):
    host = s(base_url).lower()
    return "localhost" in host or "127.0.0.1" in host


def configure_session(session, spec):
    key = resolve_api_key(spec["cfg"])
    if key:
        session.headers["Authorization"] = f"Bearer {key}"
        spec["cfg"]["_api_key_set"] = True
    else:
        spec["cfg"]["_api_key_set"] = False
    return session


def probe_endpoint(session, spec):
    cfg = spec["cfg"]
    base = s(cfg["base_url"]).rstrip("/")
    if is_local_ollama(base):
        try:
            return session.get(base + "/api/version", timeout=10).json().get("version", "unknown")
        except Exception as e:
            sys.exit(f"Cannot reach Ollama at {base}: {e}")
    if not resolve_api_key(cfg):
        sys.exit(
            f"base_url {base} is not local Ollama; set TYPESAFE_API_KEY "
            "(or Config api_key) for Authorization."
        )
    try:
        r = session.get(base + "/v1/models", timeout=15)
        if r.status_code >= 400:
            sys.exit(f"Cannot reach decision API at {base}/v1/models: HTTP {r.status_code}: {r.text[:200]}")
        return f"remote:{base}"
    except Exception as e:
        sys.exit(f"Cannot reach decision API at {base}: {e}")


def call_model(session, spec, body):
    cfg = spec["cfg"]
    url = s(cfg["base_url"]).rstrip("/") + s(cfg["endpoint"])
    last = None
    for attempt in range(cfg["retries"] + 1):
        try:
            r = session.post(url, json=body, timeout=cfg["timeout_seconds"])
            if r.status_code >= 500:
                raise RuntimeError(f"HTTP {r.status_code}: {r.text[:200]}")
            if r.status_code >= 400:
                raise ValueError(f"HTTP {r.status_code}: {r.text[:300]}")
            payload = r.json() if r.content else {}
            usage = payload.get("usage") if isinstance(payload.get("usage"), dict) else None
            return payload.get("answers", {}), None, usage
        except ValueError as e:
            return {}, str(e), None
        except Exception as e:
            last = str(e)
            time.sleep(1.5 * (attempt + 1))
    return {}, last, None


def usage_totals(results):
    inp = out = n = 0
    for obj in results.values():
        u = obj.get("usage") if isinstance(obj, dict) else None
        if not isinstance(u, dict):
            continue
        it, ot = u.get("input_tokens"), u.get("output_tokens")
        if isinstance(it, (int, float)):
            inp += int(it)
            n += 1
        if isinstance(ot, (int, float)):
            out += int(ot)
    return {"input_tokens": inp, "output_tokens": out, "records_with_usage": n,
            "total_tokens": inp + out}


def parse_answer(q, labels, ans):
    out = {"answer": "", "confidence": None, "value": None}
    if not isinstance(ans, dict):
        return out
    probs = ans.get("probabilities") if isinstance(ans.get("probabilities"), dict) else {}
    conf = ans.get("confidence")
    out["confidence"] = conf if num(conf) else None
    t = q["type"]
    if t == "choice":
        choice = ans.get("choice")
        if choice is None and probs:
            choice = max(probs, key=probs.get)
        out["answer"] = "" if choice is None else str(choice)
        p = probs.get(choice)
        out["value"] = p if num(p) else None
    elif t == "noul":
        cands = [ans.get("noul"), ans.get("probability"), ans.get("p_true"),
                 probs.get("true"), probs.get("True"), probs.get("yes")]
        p = next((c for c in cands if num(c)), None)
        out["value"] = p
        if p is not None:
            out["answer"] = "true" if p >= 0.5 else "false"
            if out["confidence"] is None:
                out["confidence"] = max(p, 1 - p)
    elif t == "score":
        val = next((ans[k] for k in ("score", "value") if num(ans.get(k))), None)
        if val is None and probs:
            val = sum(i * (probs.get(l["label"]) or 0) for i, l in enumerate(labels))
        out["value"] = val
        if val is not None and labels:
            idx = min(max(int(round(val)), 0), len(labels) - 1)
            out["answer"] = labels[idx]["label"]
    return out


def review_reason(q, parsed, threshold):
    if parsed["answer"] == "":
        return "no answer"
    thr = q["review_below"] if q["review_below"] is not None else threshold
    if q["type"] == "noul":
        p = parsed["value"]
        if p is not None and max(p, 1 - p) < thr:
            return f"{q['id']}: uncertain (p_true={p:.2f})"
    elif parsed["confidence"] is not None and parsed["confidence"] < thr:
        return f"{q['id']}: low confidence ({parsed['confidence']:.2f})"
    return ""


# ----------------------------------------------------------------- output
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
DIFF_FILL = PatternFill("solid", fgColor="FCE4D6")
BODY = Font(name="Arial", size=10)


def write_sheet(ws, headers, rows, widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def clear_engine_output_sheets(workbook_path, engine):
    """Clear only this engine's output_* sheets (+ legacy unprefixed once). Never touch the other engine."""
    wb = load_workbook(workbook_path, data_only=False)
    prefix = f"output_{engine}_"
    stale = [n for n in wb.sheetnames if n.startswith(prefix)]
    legacy = [n for n in LEGACY_OUTPUT_SHEETS if n in wb.sheetnames]
    removed = stale + legacy
    for name in removed:
        del wb[name]
    if removed:
        wb.save(workbook_path)
        print(f"Cleared sheets for engine={engine}: {', '.join(removed)}")
    else:
        print(f"No previous output_{engine}_* sheets to clear.")


def replace_sheet(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def _results_answer_map(wb, sheet_name):
    """Map (record_id, question_id) -> (answer, confidence) from an output_*_Results sheet."""
    if sheet_name not in wb.sheetnames:
        return {}
    rows = read_table(wb[sheet_name])
    out = {}
    for r in rows:
        rid = s(r.get("record_id"))
        if not rid:
            continue
        for k, v in r.items():
            if k.endswith("__answer"):
                qid = k[: -len("__answer")]
                conf = to_float(r.get(f"{qid}__confidence"))
                out[(rid, qid)] = (s(v), conf)
    return out


def build_compare_sheet(wb, spec):
    """Refresh output_compare when both engine Results sheets exist."""
    nim = _results_answer_map(wb, "output_nimble_Results")
    jev = _results_answer_map(wb, "output_jev_Results")
    if not nim or not jev:
        return False
    expected = {}
    for rec in spec["records"]:
        for qid, exp in rec["expected"].items():
            expected[(rec["id"], qid)] = exp
    keys = sorted(set(nim) | set(jev))
    rows = []
    for rid, qid in keys:
        na, nc = nim.get((rid, qid), ("", None))
        ja, jc = jev.get((rid, qid), ("", None))
        if na == "" and ja == "":
            continue
        exp = expected.get((rid, qid), "")
        agree = na != "" and ja != "" and na.lower() == ja.lower()
        rows.append([
            rid, qid, na, ja, agree, nc, jc, exp,
            (na.lower() == exp) if exp and na else "",
            (ja.lower() == exp) if exp and ja else "",
        ])
    ws = replace_sheet(wb, "output_compare")
    write_sheet(ws,
                ["record_id", "question_id", "nimble_answer", "jev_answer", "agree",
                 "nimble_confidence", "jev_confidence", "expected",
                 "nimble_correct", "jev_correct"],
                rows, [12, 22, 14, 14, 8, 14, 14, 12, 12, 12])
    for i, row in enumerate(rows, start=2):
        if row[4] is False:
            for c in ws[i]:
                c.fill = DIFF_FILL
    print("Updated output_compare (both engines present).")
    return True


def build_outputs(spec, results, run_info, workbook_path, engine):
    """Write output_{engine}_* sheets only; refresh output_compare if both engines exist."""
    names = output_sheet_names(engine)
    qmap = {q["id"]: q for q in spec["questions"]}
    done = [r for r in spec["records"] if r["id"] in results]
    asked_ids = []
    for r in done:
        for q in questions_for(r, spec):
            if q["id"] not in asked_ids:
                asked_ids.append(q["id"])
    meta_cols = list(done[0]["meta"].keys()) if done else ["record_id"]

    headers = meta_cols + ["text"]
    for qid in asked_ids:
        headers += [f"{qid}__answer", f"{qid}__confidence", f"{qid}__value"]
    headers += ["needs_review", "review_reasons", "error"]

    stats = {qid: {"asked": 0, "answered": 0, "review": 0, "conf": [], "exp": 0,
                   "ok": 0, "conf_ok": [], "conf_bad": []} for qid in asked_ids}
    rows, raw_rows, review_flags = [], [], []
    for rec in done:
        res = results[rec["id"]]
        row = [rec["meta"].get(c, "") for c in meta_cols] + [rec["text"]]
        reasons = []
        for qid in asked_ids:
            q = qmap[qid]
            applies = q in questions_for(rec, spec)
            if not applies:
                row += ["", "", ""]
                continue
            st = stats[qid]
            st["asked"] += 1
            parsed = parse_answer(q, spec["labels"].get(qid, []), res["answers"].get(qid))
            row += [parsed["answer"], parsed["confidence"], parsed["value"]]
            if parsed["answer"] != "":
                st["answered"] += 1
            if parsed["confidence"] is not None:
                st["conf"].append(parsed["confidence"])
            why = review_reason(q, parsed, spec["cfg"]["default_review_below"])
            if why:
                reasons.append(why)
                st["review"] += 1
            exp = rec["expected"].get(qid)
            if exp and parsed["answer"] != "":
                st["exp"] += 1
                correct = parsed["answer"].lower() == exp
                st["ok"] += int(correct)
                if parsed["confidence"] is not None:
                    (st["conf_ok"] if correct else st["conf_bad"]).append(parsed["confidence"])
        if res.get("error"):
            reasons.append("request failed")
        row += [bool(reasons), "; ".join(reasons), res.get("error") or ""]
        review_flags.append(bool(reasons))
        rows.append(row)
        raw_rows.append([rec["id"], json.dumps(res.get("usage") or {}), json.dumps(res["answers"])])

    wb = load_workbook(workbook_path, data_only=False)
    for name in names:
        if name in wb.sheetnames:
            del wb[name]

    ws = replace_sheet(wb, names[0])
    widths = [12] * len(meta_cols) + [60] + [16, 12, 10] * len(asked_ids) + [12, 45, 30]
    write_sheet(ws, headers, rows, widths)
    for i, flag in enumerate(review_flags, start=2):
        if flag:
            for c in ws[i]:
                c.fill = REVIEW_FILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "0.000"
        row[len(meta_cols)].alignment = Alignment(wrap_text=True, vertical="top")

    write_sheet(replace_sheet(wb, names[1]),
                ["record_id", "usage_json", "raw_answers_json"], raw_rows, [12, 40, 120])

    mean = lambda xs: round(sum(xs) / len(xs), 3) if xs else ""
    srows = []
    for qid in asked_ids:
        st = stats[qid]
        srows.append([qid, qmap[qid]["type"], st["asked"], st["answered"], st["review"],
                      mean(st["conf"]), st["exp"], st["ok"],
                      round(st["ok"] / st["exp"], 3) if st["exp"] else "",
                      mean(st["conf_ok"]), mean(st["conf_bad"])])
    write_sheet(replace_sheet(wb, names[2]),
                ["question_id", "type", "asked", "answered", "flagged_for_review", "mean_confidence",
                 "with_expected", "correct", "accuracy", "mean_conf_when_correct", "mean_conf_when_wrong"],
                srows, [24, 10, 8, 10, 18, 16, 14, 9, 10, 20, 20])

    write_sheet(replace_sheet(wb, names[3]),
                ["item", "value"], [[k, str(v)] for k, v in run_info.items()], [26, 70])

    build_compare_sheet(wb, spec)
    wb.save(workbook_path)
    return names


# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--engine", choices=sorted(ENGINES), default=None,
                    help="nimble (local Ollama) or jev (TypeSafe); overrides Config engine")
    ap.add_argument("--validate-only", action="store_true", help="check the workbook and exit")
    ap.add_argument("--test", action="store_true", help="only the first test_first_n records")
    ap.add_argument("--resume", metavar="JSONL", help="continue from a checkpoint file")
    ap.add_argument("--out-dir", default=None,
                    help="directory for the JSONL resume checkpoint "
                         "(default: same folder as the workbook)")
    args = ap.parse_args()

    workbook_path = Path(args.workbook).resolve()
    spec = load_spec(workbook_path)
    errs = validate(spec)
    if errs:
        print(f"{len(errs)} problem(s) in the workbook:")
        for e in errs:
            print("  -", e)
        sys.exit(1)

    engine = resolve_engine(spec["cfg"], args.engine)
    apply_engine_preset(spec["cfg"], engine)
    print(f"Engine: {engine} (model={spec['cfg']['model']}, base_url={spec['cfg']['base_url']})")

    qhash = questions_hash(spec)
    active = [q for q in spec["questions"] if q["active"]]
    print(f"Workbook OK: {len(active)} active question(s), {len(spec['records'])} record(s).")
    if args.validate_only:
        return

    records = spec["records"]
    if args.test and spec["cfg"]["test_first_n"] > 0:
        records = records[:spec["cfg"]["test_first_n"]]
        spec["records"] = records
        print(f"Test mode: first {len(records)} record(s).")

    cfg = spec["cfg"]
    base = s(cfg["base_url"]).rstrip("/")
    session = configure_session(requests.Session(), spec)
    endpoint_info = probe_endpoint(session, spec)
    print(f"Endpoint OK: {base} ({endpoint_info})")

    # Only clear this engine's sheets — keep the other engine's results for compare.
    clear_engine_output_sheets(workbook_path, engine)

    ckpt_dir = Path(args.out_dir) if args.out_dir else workbook_path.parent
    ckpt_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    results = {}
    if args.resume:
        ckpt = Path(args.resume)
        for line in ckpt.read_text(encoding="utf-8").splitlines():
            obj = json.loads(line)
            if "_meta" in obj:
                if obj["_meta"]["questions_hash"] != qhash:
                    sys.exit("The Questions/Labels changed since this checkpoint was written. "
                             "Start a new run instead of resuming.")
            else:
                results[obj["record_id"]] = obj
        print(f"Resuming: {len(results)} record(s) already done.")
    else:
        ckpt = ckpt_dir / f"{workbook_path.stem}_{engine}_{stamp}.jsonl"
        ckpt.write_text(json.dumps({"_meta": {"questions_hash": qhash, "workbook": str(workbook_path),
                                              "engine": engine, "started": stamp}}) + "\n",
                        encoding="utf-8")

    started = time.time()
    todo = [r for r in records if r["id"] not in results]
    try:
        with ckpt.open("a", encoding="utf-8") as f:
            for i, rec in enumerate(todo, 1):
                qs = questions_for(rec, spec)
                t0 = time.time()
                if qs:
                    answers, error, usage = call_model(session, spec, build_request(rec, qs, spec))
                else:
                    answers, error, usage = {}, "no questions apply to this equipment_type", None
                obj = {"record_id": rec["id"], "answers": answers, "error": error,
                       "usage": usage, "seconds": round(time.time() - t0, 3)}
                results[rec["id"]] = obj
                f.write(json.dumps(obj) + "\n")
                f.flush()
                u = usage or {}
                utxt = ""
                if isinstance(u.get("input_tokens"), (int, float)) or isinstance(u.get("output_tokens"), (int, float)):
                    utxt = f" in={u.get('input_tokens')} out={u.get('output_tokens')}"
                print(f"[{i}/{len(todo)}] {rec['id']} {'ERROR ' + error if error else 'ok'} "
                      f"({obj['seconds']}s){utxt}")
    except KeyboardInterrupt:
        print(f"\nInterrupted. Resume with: python {sys.argv[0]} {workbook_path} "
              f"--engine {engine} --resume {ckpt}")

    totals = usage_totals(results)
    sheet_names = output_sheet_names(engine)
    info = {"run_started": stamp, "workbook": str(workbook_path), "engine": engine,
            "model": cfg["model"], "base_url": base, "endpoint_info": endpoint_info,
            "questions_hash": qhash, "api_key_configured": bool(cfg.get("_api_key_set")),
            "records_in_run": len(records),
            "records_completed": sum(1 for r in records if r["id"] in results),
            "failed_records": sum(1 for r in records if r["id"] in results and results[r["id"]].get("error")),
            "review_threshold_default": cfg["default_review_below"],
            "elapsed_seconds": round(time.time() - started, 1), "checkpoint": str(ckpt),
            "input_tokens": totals["input_tokens"], "output_tokens": totals["output_tokens"],
            "total_tokens": totals["total_tokens"],
            "records_with_usage": totals["records_with_usage"],
            "output_sheets": ", ".join(sheet_names)}
    written = build_outputs(spec, results, info, workbook_path, engine)
    print(f"Updated {workbook_path} with sheets: {', '.join(written)}")
    print(f"Tokens: input={totals['input_tokens']} output={totals['output_tokens']} "
          f"total={totals['total_tokens']} (from {totals['records_with_usage']} record(s))")
    print(f"Checkpoint (for --resume): {ckpt}")


if __name__ == "__main__":
    main()
